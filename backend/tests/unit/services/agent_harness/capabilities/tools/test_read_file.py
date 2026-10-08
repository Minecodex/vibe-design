from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook
from fastapi import HTTPException

from app.api.v1.endpoints import harness as harness_endpoint
from app.services.agent_harness.capabilities.tools.read_file import ReadFileInput, ReadFileTool
from app.services.agent_harness.core.context import HarnessContext


def _ctx(tmp_path: Path) -> HarnessContext:
    ctx = HarnessContext(
        user_id=1,
        conversation_id="conv-harness-files",
        run_id="run-harness-files",
        workspace_root=tmp_path,
    )
    ctx.ensure_dirs()
    return ctx


def _write_large_text_file(path: Path) -> None:
    line = "x" * 100
    path.write_text("\n".join(f"{line}{index:05d}" for index in range(55_000)), encoding="utf-8")


@pytest.mark.asyncio
async def test_read_file_reads_small_plain_text_to_eof_with_line_numbers(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.project_dir / "notes.txt"
    path.write_text("alpha\nbeta", encoding="utf-8")

    result = await ReadFileTool().execute(ReadFileInput(base="work", file_path="notes.txt"), ctx)

    assert result.is_error is False
    assert result.output == "     1\talpha\n     2\tbeta"
    assert result.metadata["truncated_by_window"] is False
    assert result.metadata["has_more"] is False


@pytest.mark.asyncio
async def test_read_file_can_page_plain_text_references(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.reference_inputs_dir / "notes.txt"
    path.write_text("line 1\nline 2\nline 3\nline 4", encoding="utf-8")

    result = await ReadFileTool().execute(
        ReadFileInput(base="references", file_path="inputs/notes.txt", offset=1, limit=2),
        ctx,
    )

    assert "lines 2-3 of 4" in result.output
    assert "continue with offset=3" in result.output
    assert result.output.endswith("     2\tline 2\n     3\tline 3")
    assert result.metadata["path"] == "references/inputs/notes.txt"
    assert result.metadata["start_line"] == 2
    assert result.metadata["num_lines"] == 2
    assert result.metadata["total_lines"] == 4
    assert result.metadata["truncated_by_window"] is True


@pytest.mark.asyncio
async def test_read_file_can_read_docx_references(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.reference_inputs_dir / "brief.docx"
    document = Document()
    document.add_paragraph("Project brief")
    document.add_paragraph("Second paragraph")
    document.save(path)

    result = await ReadFileTool().execute(
        ReadFileInput(base="references", file_path="inputs/brief.docx"),
        ctx,
    )

    assert result.output == "     1\tProject brief\n     2\tSecond paragraph"
    assert result.metadata["path"] == "references/inputs/brief.docx"
    assert result.metadata["start_line"] == 1
    assert result.metadata["truncated_by_window"] is False


@pytest.mark.asyncio
async def test_read_file_directory_path_points_agent_to_list_files(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    (ctx.skill_dir / "scripts" / "verify-output").mkdir(parents=True)
    (ctx.skill_dir / "scripts" / "verify-output" / "example.png").write_bytes(b"png")

    result = await ReadFileTool().execute(
        ReadFileInput(base="skill", file_path="scripts/verify-output"),
        ctx,
    )

    assert result.is_error is True
    assert result.metadata["failure_kind"] == "path_is_directory"
    assert "directory" in result.output
    assert 'list_files(path="skill/scripts/verify-output")' in result.output


@pytest.mark.asyncio
async def test_read_file_can_read_xlsx_references_with_windowing(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.reference_inputs_dir / "budget.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet["A1"] = "Item"
    sheet["B1"] = "Amount"
    sheet["A2"] = "Tea"
    sheet["B2"] = 18
    workbook.save(path)

    result = await ReadFileTool().execute(
        ReadFileInput(base="references", file_path="inputs/budget.xlsx", offset=1, limit=1),
        ctx,
    )

    assert "lines 2-2 of 3" in result.output
    assert result.output.endswith("     2\t1\tItem\tAmount")
    assert result.metadata["path"] == "references/inputs/budget.xlsx"
    assert result.metadata["start_line"] == 2
    assert result.metadata["num_lines"] == 1
    assert result.metadata["total_lines"] == 3
    assert result.metadata["truncated_by_window"] is True
    assert result.metadata["has_more"] is True
    assert result.metadata["next_offset"] == 2


@pytest.mark.asyncio
async def test_read_file_repeated_unchanged_window_returns_compact_result(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.project_dir / "notes.txt"
    path.write_text("line 1\nline 2\nline 3", encoding="utf-8")
    tool = ReadFileTool()
    params = ReadFileInput(base="work", file_path="notes.txt", offset=0, limit=2)

    first = await tool.execute(params, ctx)
    second = await tool.execute(params, ctx)

    assert "line 1" in first.output
    assert second.metadata["unchanged"] is True
    assert second.metadata["repeated_read"] is True
    assert "unchanged" in second.output.lower()
    assert "line 1" not in second.output


@pytest.mark.asyncio
async def test_read_file_repeated_unchanged_window_keeps_returning_stub(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.project_dir / "notes.txt"
    path.write_text("line 1\nline 2\nline 3", encoding="utf-8")
    tool = ReadFileTool()
    params = ReadFileInput(base="work", file_path="notes.txt", offset=0, limit=2)

    first = await tool.execute(params, ctx)
    second = await tool.execute(params, ctx)
    third = await tool.execute(params, ctx)

    assert first.is_error is False
    assert second.is_error is False
    assert third.is_error is False
    assert third.metadata["unchanged"] is True
    assert "File unchanged since last read" in third.output


@pytest.mark.asyncio
async def test_read_file_repeated_unchanged_stub_allows_different_window(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.project_dir / "notes.txt"
    path.write_text("line 1\nline 2\nline 3", encoding="utf-8")
    tool = ReadFileTool()

    await tool.execute(ReadFileInput(base="work", file_path="notes.txt", offset=0, limit=1), ctx)
    await tool.execute(ReadFileInput(base="work", file_path="notes.txt", offset=0, limit=1), ctx)
    result = await tool.execute(ReadFileInput(base="work", file_path="notes.txt", offset=1, limit=1), ctx)

    assert result.is_error is False
    assert result.metadata["start_line"] == 2
    assert "     2\tline 2" in result.output


@pytest.mark.asyncio
async def test_read_file_changed_file_invalidates_compact_unchanged_result(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.project_dir / "notes.txt"
    path.write_text("line 1\nline 2", encoding="utf-8")
    tool = ReadFileTool()
    params = ReadFileInput(base="work", file_path="notes.txt", offset=0, limit=2)

    await tool.execute(params, ctx)
    path.write_text("changed\nline 2", encoding="utf-8")
    result = await tool.execute(params, ctx)

    assert result.metadata["unchanged"] is False
    assert "     1\tchanged" in result.output


@pytest.mark.asyncio
async def test_read_file_can_use_doc_fallback_extractor(tmp_path: Path, monkeypatch) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.reference_inputs_dir / "legacy.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy-doc-placeholder")

    # Exercise the fallback independently of the installed MarkItDown version.
    monkeypatch.setattr(
        "app.services.agent_harness.capabilities.tools._internal.harness_file_readers._extract_with_markitdown",
        lambda _resolved: None,
    )

    monkeypatch.setattr(
        "app.services.agent_harness.capabilities.tools._internal.harness_file_readers._extract_doc_text_with_soffice",
        lambda resolved: "Legacy doc content",
    )

    result = await ReadFileTool().execute(
        ReadFileInput(base="references", file_path="inputs/legacy.doc"),
        ctx,
    )

    assert result.output == "     1\tLegacy doc content"
    assert result.metadata["path"] == "references/inputs/legacy.doc"


@pytest.mark.asyncio
async def test_read_file_large_file_without_limit_asks_for_window_or_search(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.project_dir / "big.log"
    _write_large_text_file(path)

    result = await ReadFileTool().execute(ReadFileInput(base="work", file_path="big.log"), ctx)

    assert result.is_error is True
    assert result.metadata["failure_kind"] == "file_read_limit"
    assert "Use offset and limit" in result.output
    assert "grep_files" in result.output


@pytest.mark.asyncio
async def test_read_file_large_file_with_limit_streams_requested_window(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.project_dir / "big.log"
    _write_large_text_file(path)

    result = await ReadFileTool().execute(
        ReadFileInput(base="work", file_path="big.log", offset=50_000, limit=2),
        ctx,
    )

    assert result.is_error is False
    assert "lines 50001-50002 of 55000" in result.output
    assert "continue with offset=50002" in result.output
    assert " 50001\t" in result.output
    assert result.metadata["start_line"] == 50_001
    assert result.metadata["num_lines"] == 2
    assert result.metadata["total_lines"] == 55_000
    assert result.metadata["has_more"] is True
    assert result.metadata["next_offset"] == 50_002


@pytest.mark.asyncio
async def test_read_file_output_over_budget_requires_smaller_window(tmp_path: Path) -> None:
    ctx = _ctx(tmp_path)
    path = ctx.project_dir / "notes.txt"
    path.write_text("\n".join(f"line {index}" for index in range(20)), encoding="utf-8")
    tool = ReadFileTool()
    tool.max_result_size_chars = 40

    result = await tool.execute(ReadFileInput(base="work", file_path="notes.txt"), ctx)

    assert result.is_error is True
    assert result.metadata["failure_kind"] == "file_read_output_too_large"
    assert "Retry with a smaller limit" in result.output


@pytest.mark.asyncio
async def test_upload_harness_attachment_rejects_unsupported_files(monkeypatch):
    class DummyUploadFile:
        def __init__(self, filename: str, content: bytes, content_type: str) -> None:
            self.filename = filename
            self.content_type = content_type
            self._content = content
            self._offset = 0

        async def read(self, size: int = -1) -> bytes:
            if self._offset >= len(self._content):
                return b""
            if size < 0:
                size = len(self._content) - self._offset
            chunk = self._content[self._offset:self._offset + size]
            self._offset += len(chunk)
            return chunk

    user = SimpleNamespace(id=7)
    monkeypatch.setattr(
        "app.services.agent_harness.workspace.conversation.conversation_meta_store.get_conversation",
        lambda user_id, conversation_id: {"id": conversation_id},
    )
    monkeypatch.setattr(
        "app.services.agent_harness.workspace.generated_content.asset_store.register_asset_bytes",
        lambda *args, **kwargs: pytest.fail("register_asset_bytes should not be called"),
    )

    with pytest.raises(HTTPException) as exc:
        await harness_endpoint.upload_harness_attachment(
            conversation_id="conv-1",
            user=user,
            file=DummyUploadFile("archive.zip", b"zip-bytes", "application/zip"),
        )

    assert exc.value.status_code == 400
    assert "Unsupported file format" in str(exc.value.detail)


@pytest.mark.asyncio
async def test_upload_harness_attachment_rejects_oversized_file_without_unbounded_read(monkeypatch, tmp_path):
    class ChunkedUploadFile:
        def __init__(self, filename: str, content: bytes, content_type: str) -> None:
            self.filename = filename
            self.content_type = content_type
            self._content = content
            self._offset = 0
            self.read_sizes: list[int] = []

        async def read(self, size: int = -1) -> bytes:
            if size < 0:
                raise AssertionError("upload endpoint must not perform an unbounded read")
            self.read_sizes.append(size)
            if self._offset >= len(self._content):
                return b""
            chunk = self._content[self._offset:self._offset + min(size, 3)]
            self._offset += len(chunk)
            return chunk

    user = SimpleNamespace(id=7)
    conversation_root = tmp_path / "users" / "7" / "conversations" / "conv-1"
    monkeypatch.setattr("app.core.config.settings.HARNESS_ATTACHMENT_UPLOAD_MAX_BYTES", 5, raising=False)
    monkeypatch.setattr(
        "app.services.agent_harness.workspace.conversation.conversation_meta_store.get_conversation",
        lambda user_id, conversation_id: {"id": conversation_id},
    )
    monkeypatch.setattr(
        "app.services.agent_harness.workspace.conversation.conversation_meta_store.get_conversation_dir",
        lambda user_id, conversation_id: conversation_root,
    )
    monkeypatch.setattr(
        "app.services.agent_harness.workspace.generated_content.asset_store.register_asset",
        lambda *args, **kwargs: pytest.fail("register_asset should not be called"),
    )
    upload = ChunkedUploadFile("brief.txt", b"123456", "text/plain")

    with pytest.raises(HTTPException) as exc:
        await harness_endpoint.upload_harness_attachment(
            conversation_id="conv-1",
            user=user,
            file=upload,
        )

    assert exc.value.status_code == 413
    assert upload.read_sizes


@pytest.mark.asyncio
async def test_upload_harness_attachment_accepts_text_and_office_formats(monkeypatch, tmp_path):
    class DummyUploadFile:
        def __init__(self, filename: str, content: bytes, content_type: str) -> None:
            self.filename = filename
            self.content_type = content_type
            self._content = content
            self._offset = 0

        async def read(self, size: int = -1) -> bytes:
            if self._offset >= len(self._content):
                return b""
            if size < 0:
                size = len(self._content) - self._offset
            chunk = self._content[self._offset:self._offset + size]
            self._offset += len(chunk)
            return chunk

    user = SimpleNamespace(id=7)
    monkeypatch.setattr(
        "app.services.agent_harness.workspace.conversation.conversation_meta_store.get_conversation",
        lambda user_id, conversation_id: {"id": conversation_id},
    )
    monkeypatch.setattr(
        "app.services.agent_harness.workspace.conversation.conversation_meta_store.get_conversation_dir",
        lambda user_id, conversation_id: tmp_path / "users" / str(user_id) / "conversations" / conversation_id,
    )
    monkeypatch.setattr(
        "app.services.agent_harness.workspace.generated_content.asset_store.register_asset",
        lambda *args, **kwargs: {
            "path": "references/inputs/upload_001/source.docx",
            "asset_id": "upload_001",
            "size": 10,
        },
    )

    response = await harness_endpoint.upload_harness_attachment(
        conversation_id="conv-1",
        user=user,
        file=DummyUploadFile(
            "brief.docx",
            b"docx-bytes",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
    )

    assert response["type"] == "file"
    assert response["url"] == "references/inputs/upload_001/source.docx"
    assert response["filename"] == "brief.docx"

