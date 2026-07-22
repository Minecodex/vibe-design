from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from app.services.agent_harness.isolation.security.paths import MAX_TEXT_FILE_READ_BYTES, read_text_file_window

TEXTUAL_UPLOAD_EXTENSIONS = frozenset({
    "txt",
    "md",
    "markdown",
    "csv",
    "json",
    "jsonl",
    "yaml",
    "yml",
    "toml",
    "xml",
    "html",
    "htm",
    "css",
    "js",
    "ts",
    "jsx",
    "tsx",
    "py",
    "sh",
    "sql",
    "ini",
    "conf",
    "log",
    "docx",
    "doc",
    "xlsx",
})

UPLOAD_IMAGE_MIME_PREFIX = "image/"
UPLOAD_TEXT_MIME_PREFIX = "text/"
UPLOAD_ALLOWED_MIME_TYPES = frozenset({
    "application/json",
    "application/xml",
    "application/xhtml+xml",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
})

OFFICE_READ_EXTENSIONS = frozenset({"doc", "docx", "xlsx"})


def is_supported_harness_upload(name: str, mime_type: str | None = None) -> bool:
    extension = _get_extension(name)
    normalized_mime_type = str(mime_type or "").strip().lower()
    if extension in TEXTUAL_UPLOAD_EXTENSIONS:
        return True
    if normalized_mime_type.startswith(UPLOAD_IMAGE_MIME_PREFIX):
        return True
    if normalized_mime_type.startswith(UPLOAD_TEXT_MIME_PREFIX):
        return True
    return normalized_mime_type in UPLOAD_ALLOWED_MIME_TYPES


def read_harness_file_window(
    resolved: Path,
    *,
    offset: int = 0,
    limit: int | None = None,
    max_size_bytes: int = MAX_TEXT_FILE_READ_BYTES,
) -> dict[str, object]:
    extension = _get_extension(resolved.name)
    if extension in OFFICE_READ_EXTENSIONS:
        content = extract_harness_file_text(resolved)
        return _window_text_content(
            content,
            source_size_bytes=resolved.stat().st_size,
            offset=offset,
            limit=limit,
            max_size_bytes=max_size_bytes,
        )

    return read_text_file_window(
        resolved,
        offset=offset,
        limit=limit,
        max_size_bytes=max_size_bytes,
    )


def extract_harness_file_text(resolved: Path) -> str:
    extension = _get_extension(resolved.name)
    if extension == "docx":
        extracted = _extract_with_markitdown(resolved)
        if extracted:
            return extracted
        return _extract_docx_text(resolved)
    if extension == "doc":
        extracted = _extract_with_markitdown(resolved)
        if extracted:
            return extracted
        return _extract_doc_text_with_soffice(resolved)
    if extension == "xlsx":
        return _extract_xlsx_text(resolved)
    raise ValueError(f"Unsupported file type: {resolved.name}")


def _extract_with_markitdown(resolved: Path) -> str | None:
    try:
        from markitdown import MarkItDown
    except ImportError:
        return None

    try:
        converter = MarkItDown()
        result = converter.convert(str(resolved))
    except Exception:
        return None

    for candidate in (
        getattr(result, "text_content", None),
        getattr(result, "text", None),
        getattr(result, "markdown", None),
    ):
        text = str(candidate or "").strip()
        if text:
            return text
    return None


def _extract_docx_text(resolved: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - declared in backend/requirements.txt
        raise ValueError("python-docx is required to read DOCX files") from exc

    document = Document(str(resolved))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            values = [value for value in values if value]
            if values:
                rows.append("\t".join(values))
        if rows:
            parts.append(f"[Table {table_index}]")
            parts.extend(rows)

    return "\n".join(parts).strip()


def _extract_doc_text_with_soffice(resolved: Path) -> str:
    soffice_bin = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice_bin is None:
        raise ValueError("LibreOffice soffice is required to read DOC files")

    with tempfile.TemporaryDirectory(prefix="harness-doc-read-") as tmp_dir:
        outdir = Path(tmp_dir)
        command = [
            soffice_bin,
            "--headless",
            "--convert-to",
            "txt:Text",
            "--outdir",
            str(outdir),
            str(resolved),
        ]
        result = subprocess.run(command, capture_output=True, text=True)
        if result.returncode != 0:
            raise ValueError((result.stderr or result.stdout or "Failed to read DOC file").strip())

        candidates = list(outdir.glob(f"{resolved.stem}*.txt"))
        if not candidates:
            raise ValueError("LibreOffice did not produce a readable text file")

        return candidates[0].read_text(encoding="utf-8", errors="replace").strip()


def _extract_xlsx_text(resolved: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - declared in backend/requirements.txt
        raise ValueError("openpyxl is required to read XLSX files") from exc

    workbook = load_workbook(resolved, data_only=False)
    parts: list[str] = []

    for sheet in workbook.worksheets:
        parts.append(f"[Sheet] {sheet.title}")
        max_row = sheet.max_row or 0
        max_col = sheet.max_column or 0
        if max_row <= 0 or max_col <= 0:
            continue

        for row_index, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True),
            start=1,
        ):
            values = [_format_cell_value(value) for value in row]
            while values and not values[-1]:
                values.pop()
            if not values:
                continue
            parts.append(f"{row_index}\t" + "\t".join(values))
        parts.append("")

    return "\n".join(parts).strip()


def _format_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return str(value)


def _window_text_content(
    content: str,
    *,
    source_size_bytes: int,
    offset: int = 0,
    limit: int | None = None,
    max_size_bytes: int = MAX_TEXT_FILE_READ_BYTES,
) -> dict[str, object]:
    if source_size_bytes > max_size_bytes:
        raise ValueError(f"File is too large ({source_size_bytes} bytes, max {max_size_bytes} bytes)")

    lines = content.splitlines()
    start_index = min(max(offset, 0), len(lines))
    if limit is None:
        end_index = len(lines)
    else:
        end_index = min(start_index + max(limit, 0), len(lines))
    selected = "\n".join(lines[start_index:end_index])
    return {
        "content": selected,
        "start_line": start_index + 1,
        "num_lines": end_index - start_index,
        "total_lines": len(lines),
        "truncated_by_window": start_index > 0 or end_index < len(lines),
        "size_bytes": source_size_bytes,
        "max_read_size_bytes": max_size_bytes,
    }


def _get_extension(name: str) -> str:
    return Path(str(name or "")).suffix.lower().lstrip(".")
