"""Temporary Office session handling for homepage Univer integration."""

from __future__ import annotations

import csv
import io
import base64
import hashlib
import html
import json
import logging
import os
import shutil
import subprocess
import tempfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
import time
from typing import Literal
from uuid import uuid4
from zipfile import ZipFile

from lxml import etree

from app.core.config import settings
from .conversation_meta_store import get_conversation_dir
from .workspace_preview_service import get_workspace_file_path
from app.services.agent_harness.capabilities.skills.docx.scripts.office.soffice import get_soffice_env

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
except ImportError:  # pragma: no cover - dependency declared in requirements.txt
    Document = None
    WD_ALIGN_PARAGRAPH = None
    CT_Tbl = None
    CT_P = None
    Table = None
    _Cell = None
    Paragraph = None
    Run = None

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - dependency declared in requirements.txt
    Workbook = None
    load_workbook = None
    Alignment = None
    Border = None
    Font = None
    PatternFill = None
    Side = None
    get_column_letter = None

OfficeFileKind = Literal["doc", "sheet", "presentation"]

_OFFICE_SESSIONS: dict[str, dict] = {}
logger = logging.getLogger(__name__)

_DEFAULT_DOCUMENT_STYLE = {
    "pageSize": {
        "width": 595 / 0.75,
        "height": 842 / 0.75,
    },
    "documentFlavor": 1,
    "marginTop": 50,
    "marginBottom": 50,
    "marginRight": 50,
    "marginLeft": 50,
    "renderConfig": {
        "zeroWidthParagraphBreak": 0,
        "vertexAngle": 0,
        "centerAngle": 0,
        "background": {"rgb": "#ccc"},
    },
    "autoHyphenation": 1,
    "doNotHyphenateCaps": 0,
    "consecutiveHyphenLimit": 2,
    "defaultHeaderId": "",
    "defaultFooterId": "",
    "evenPageHeaderId": "",
    "evenPageFooterId": "",
    "firstPageHeaderId": "",
    "firstPageFooterId": "",
    "evenAndOddHeaders": 0,
    "useFirstPageHeaderFooter": 0,
    "marginHeader": 30,
    "marginFooter": 30,
}

_DEFAULT_PARAGRAPH_STYLE = {
    "spaceAbove": {"v": 5},
    "lineSpacing": 1,
    "spaceBelow": {"v": 0},
}

_WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_WORD_NAMESPACES = {"w": _WORD_NAMESPACE}
_BULLET_LIST_TYPE = "BULLET_LIST"
_ORDER_LIST_TYPE = "ORDER_LIST"
_TABLE_START = "\u001A"
_TABLE_ROW_START = "\u001B"
_TABLE_CELL_START = "\u001C"
_TABLE_CELL_END = "\u001D"
_TABLE_ROW_END = "\u000E"
_TABLE_END = "\u000F"
_CUSTOM_BLOCK = "\b"
_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PAGE_BREAK = "\f"
_WORD_HIGHLIGHT_COLORS = {
    "yellow": "#FFFF00",
    "green": "#00FF00",
    "cyan": "#00FFFF",
    "magenta": "#FF00FF",
    "blue": "#0000FF",
    "red": "#FF0000",
    "darkBlue": "#000080",
    "darkCyan": "#008080",
    "darkGreen": "#008000",
    "darkMagenta": "#800080",
    "darkRed": "#800000",
    "darkYellow": "#808000",
    "darkGray": "#808080",
    "lightGray": "#C0C0C0",
    "black": "#000000",
    "white": "#FFFFFF",
}


_XLSX_UNSUPPORTED_FEATURE_ORDER = (
    "comments",
    "conditional_formatting",
    "data_validation",
    "charts",
    "images",
    "shapes",
    "macros",
    "rich_text",
)

_XLSX_UNSUPPORTED_FEATURE_WARNINGS = {
    "comments": "检测到批注，但预览暂不支持显示。",
    "conditional_formatting": "预览暂不支持条件格式规则。",
    "data_validation": "预览暂不支持数据验证规则。",
    "charts": "检测到图表、图片或形状，但预览暂不支持显示。",
    "images": "检测到图表、图片或形状，但预览暂不支持显示。",
    "shapes": "检测到图表、图片或形状，但预览暂不支持显示。",
    "macros": "检测到宏/VBA 内容，但预览暂不支持运行或显示。",
    "rich_text": "检测到单元格内的富文本，预览时局部样式可能会丢失。",
}
_XLSX_HTML_PREVIEW_FEATURES = {"charts", "images", "shapes"}
_XLSX_HTML_PREVIEW_WARNING = "检测到图表、图片或形状，已使用 HTML 预览渲染。"
_XLSX_HTML_PREVIEW_FALLBACK_WARNING = "检测到图表、图片或形状，但无法生成 HTML 预览，已回退为表格预览。"

class _ListDefinitionInfo(dict):
    pass


class _BlockItemInfo(dict):
    pass


class _ParagraphSegmentInfo(dict):
    pass


class _NoteReferenceInfo(dict):
    pass


def _get_file_extension(file_path: str) -> str:
    return str(file_path or "").split(".")[-1].lower()


def _get_file_kind(file_path: str) -> OfficeFileKind:
    extension = _get_file_extension(file_path)
    if extension in {"doc", "docx"}:
        return "doc"
    if extension in {"xls", "xlsx", "csv"}:
        return "sheet"
    if extension in {"ppt", "pptx"}:
        return "presentation"
    raise ValueError("Unsupported office file type")


def _office_session_dir(user_id: int, conversation_id: str) -> Path:
    return get_conversation_dir(user_id, conversation_id) / ".agent" / "office_sessions"


def _office_session_record_path(user_id: int, conversation_id: str, session_id: str) -> Path:
    return _office_session_dir(user_id, conversation_id) / f"{session_id}.json"


def _build_office_session_record(
    *,
    session_id: str,
    user_id: int,
    conversation_id: str,
    file_path: str,
    file_kind: OfficeFileKind,
    file_extension: str,
    engine: str,
    readonly: bool,
    opened_at: str,
) -> dict[str, object]:
    return {
        "session_id": session_id,
        "user_id": int(user_id),
        "conversation_id": str(conversation_id),
        "file_path": str(file_path),
        "file_kind": str(file_kind),
        "file_extension": str(file_extension),
        "engine": str(engine),
        "readonly": bool(readonly),
        "opened_at": str(opened_at),
    }


def _persist_office_session_record(user_id: int, conversation_id: str, session_id: str, record: dict[str, object]) -> None:
    session_dir = _office_session_dir(user_id, conversation_id)
    session_dir.mkdir(parents=True, exist_ok=True)
    _office_session_record_path(user_id, conversation_id, session_id).write_text(
        json.dumps(record, ensure_ascii=False),
        encoding="utf-8",
    )


def _load_office_session_record(user_id: int, conversation_id: str, session_id: str) -> dict[str, object] | None:
    record_path = _office_session_record_path(user_id, conversation_id, session_id)
    if not record_path.exists():
        return None
    try:
        payload = json.loads(record_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        logger.warning(
            "Failed to read office session metadata user_id=%s conversation_id=%s session_id=%s",
            user_id,
            conversation_id,
            session_id,
        )
        return None
    return payload if isinstance(payload, dict) else None


def _delete_office_session_record(user_id: int, conversation_id: str, session_id: str) -> None:
    _office_session_record_path(user_id, conversation_id, session_id).unlink(missing_ok=True)


def _parse_office_session_opened_at(value: object) -> datetime | None:
    if not isinstance(value, str) or not value.strip():
        return None
    try:
        opened_at = datetime.fromisoformat(value)
    except ValueError:
        return None
    if opened_at.tzinfo is None:
        opened_at = opened_at.replace(tzinfo=timezone.utc)
    return opened_at.astimezone(timezone.utc)


def cleanup_expired_workspace_office_sessions(
    user_id: int,
    conversation_id: str,
    *,
    ttl_seconds: int | None = None,
    now: datetime | None = None,
) -> int:
    session_dir = _office_session_dir(user_id, conversation_id)
    if not session_dir.exists():
        return 0

    ttl = max(int(ttl_seconds if ttl_seconds is not None else settings.OFFICE_SESSION_TTL_SECONDS), 1)
    cutoff = (now or datetime.now(timezone.utc)).astimezone(timezone.utc)
    removed = 0
    for record_path in session_dir.glob("*.json"):
        try:
            payload = json.loads(record_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            logger.warning("Removing unreadable office session metadata: %s", record_path)
            payload = {}

        opened_at = _parse_office_session_opened_at(payload.get("opened_at") if isinstance(payload, dict) else None)
        if opened_at is None:
            try:
                opened_at = datetime.fromtimestamp(record_path.stat().st_mtime, timezone.utc)
            except OSError:
                opened_at = cutoff
        if (cutoff - opened_at).total_seconds() <= ttl:
            continue

        session_id = record_path.stem
        try:
            record_path.unlink(missing_ok=True)
        except OSError:
            logger.warning("Failed to remove expired office session metadata: %s", record_path)
            continue
        _OFFICE_SESSIONS.pop(session_id, None)
        removed += 1
    return removed


def _is_readonly_extension(extension: str, file_kind: OfficeFileKind) -> bool:
    if file_kind == "presentation":
        return True
    return extension in {"doc", "xls"}


def _trim_sheet_cells(cells: list[list[str | int | float | bool | None]]) -> list[list[str | int | float | bool | None]]:
    trimmed_rows: list[list[str | int | float | bool | None]] = []
    for row in cells:
        trimmed_row = list(row)
        while trimmed_row and trimmed_row[-1] is None:
            trimmed_row.pop()
        trimmed_rows.append(trimmed_row)
    while trimmed_rows and not trimmed_rows[-1]:
        trimmed_rows.pop()
    return trimmed_rows or [[]]


def _build_sheet_cell_snapshot(value, number_format: str | None = None) -> dict | None:
    if value is None:
        return None

    snapshot: dict[str, str | int | float | bool | None] = {}
    if isinstance(value, bool):
        snapshot["v"] = value
        snapshot["t"] = "b"
    elif isinstance(value, (int, float)):
        snapshot["v"] = value
        snapshot["t"] = "n"
    else:
        snapshot["v"] = str(value)
        snapshot["t"] = "s"

    if number_format:
        snapshot["numFmt"] = str(number_format)
    return snapshot


def _build_sparse_sheet_snapshot(sheet_name: str, rows: list[list[str | int | float | bool | None]]) -> dict:
    cells: dict[str, dict] = {}
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if value is None:
                continue
            cell_snapshot = _build_sheet_cell_snapshot(value)
            if cell_snapshot is not None:
                cells[f"{row_index}:{column_index}"] = cell_snapshot

    return {
        "name": sheet_name,
        "rows": {},
        "cols": {},
        "cells": cells,
        "merges": [],
        "freeze": None,
    }


def _parse_freeze_panes(freeze_panes) -> dict | None:
    if freeze_panes is None:
        return None

    coordinate = getattr(freeze_panes, "coordinate", freeze_panes)
    if not coordinate or not isinstance(coordinate, str):
        return None

    letters = "".join(character for character in coordinate if character.isalpha())
    digits = "".join(character for character in coordinate if character.isdigit())
    if not letters and not digits:
        return None

    column_index = 0
    for character in letters.upper():
        column_index = (column_index * 26) + (ord(character) - 64)

    row_index = int(digits) if digits else 1
    col_split = max(column_index - 1, 0)
    row_split = max(row_index - 1, 0)
    if col_split <= 0 and row_split <= 0:
        return None
    return {
        "rowSplit": row_split,
        "colSplit": col_split,
    }


def _build_workbook_sheet_snapshot(worksheet) -> dict:
    rows: dict[str, dict] = {}
    cols: dict[str, dict] = {}
    cells: dict[str, dict] = {}

    for row_index, row_dimension in worksheet.row_dimensions.items():
        if row_dimension.height is not None:
            rows[str(int(row_index) - 1)] = {"h": float(row_dimension.height)}

    for column_label, column_dimension in worksheet.column_dimensions.items():
        if column_dimension.width is None:
            continue
        column_index = 0
        for character in str(column_label).upper():
            if not character.isalpha():
                column_index = 0
                break
            column_index = (column_index * 26) + (ord(character) - 64)
        if column_index > 0:
            cols[str(column_index - 1)] = {"w": float(column_dimension.width)}

    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            if value is None:
                continue
            cell_snapshot = _build_sheet_cell_snapshot(value, getattr(cell, "number_format", None))
            if isinstance(value, str) and value.startswith("="):
                cell_snapshot = {
                    "v": value,
                    "t": "f",
                    "f": value[1:],
                }
                if getattr(cell, "number_format", None):
                    cell_snapshot["numFmt"] = str(cell.number_format)
            elif getattr(cell, "data_type", None) == "f":
                formula = str(value)
                cell_snapshot = {
                    "v": formula,
                    "t": "f",
                    "f": formula[1:] if formula.startswith("=") else formula,
                }
                if getattr(cell, "number_format", None):
                    cell_snapshot["numFmt"] = str(cell.number_format)
            if cell_snapshot is not None:
                cells[f"{cell.row - 1}:{cell.column - 1}"] = cell_snapshot

    merges = [
        {
            "startRow": merge_range.min_row - 1,
            "startCol": merge_range.min_col - 1,
            "endRow": merge_range.max_row - 1,
            "endCol": merge_range.max_col - 1,
        }
        for merge_range in worksheet.merged_cells.ranges
    ]

    return {
        "id": f"sheet-{worksheet.title}",
        "name": worksheet.title,
        "rows": rows,
        "cols": cols,
        "cells": cells,
        "merges": merges,
        "freeze": _parse_freeze_panes(worksheet.freeze_panes),
    }


def _copy_default_document_style() -> dict:
    return {
        **_DEFAULT_DOCUMENT_STYLE,
        "pageSize": dict(_DEFAULT_DOCUMENT_STYLE["pageSize"]),
        "renderConfig": {
            **_DEFAULT_DOCUMENT_STYLE["renderConfig"],
            "background": dict(_DEFAULT_DOCUMENT_STYLE["renderConfig"]["background"]),
        },
    }


def _copy_document_style(document_style: dict) -> dict:
    copied = dict(document_style)
    page_size = copied.get("pageSize") or {}
    render_config = copied.get("renderConfig") or {}
    copied["pageSize"] = dict(page_size)
    copied["renderConfig"] = {
        **render_config,
        "background": dict((render_config.get("background") or {})),
    }
    return copied


def _copy_default_paragraph_style() -> dict:
    return {
        **_DEFAULT_PARAGRAPH_STYLE,
        "spaceAbove": dict(_DEFAULT_PARAGRAPH_STYLE["spaceAbove"]),
        "spaceBelow": dict(_DEFAULT_PARAGRAPH_STYLE["spaceBelow"]),
    }


def _normalize_run_text(value: str | None) -> str:
    normalized = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    return normalized.replace("\n", "\v")


def _to_points(value) -> float | None:
    if value is None:
        return None
    try:
        return float(value.pt)
    except Exception:
        return None


def _points_to_pixels(value: float | None) -> float | None:
    if value is None:
        return None
    return float(value) / 0.75


def _to_pixels(value) -> float | None:
    return _points_to_pixels(_to_points(value))


def _build_number_unit(value: float | None) -> dict | None:
    if value is None:
        return None
    return {"v": float(value)}


def _map_paragraph_alignment(paragraph) -> int | None:
    alignment = getattr(paragraph, "alignment", None)
    if alignment is None:
        return None

    alignment_value = getattr(alignment, "value", alignment)
    if WD_ALIGN_PARAGRAPH is not None:
        mapping = {
            WD_ALIGN_PARAGRAPH.LEFT: 1,
            WD_ALIGN_PARAGRAPH.CENTER: 2,
            WD_ALIGN_PARAGRAPH.RIGHT: 3,
            WD_ALIGN_PARAGRAPH.JUSTIFY: 4,
            WD_ALIGN_PARAGRAPH.JUSTIFY_HI: 4,
            WD_ALIGN_PARAGRAPH.JUSTIFY_MED: 4,
            WD_ALIGN_PARAGRAPH.JUSTIFY_LOW: 4,
            WD_ALIGN_PARAGRAPH.DISTRIBUTE: 6,
        }
        for docx_alignment, univer_alignment in mapping.items():
            if alignment_value == getattr(docx_alignment, "value", docx_alignment):
                return univer_alignment

    return None


def _map_named_style_type(paragraph) -> int | None:
    style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "").strip().lower()
    mapping = {
        "normal": 1,
        "title": 2,
        "subtitle": 3,
        "heading 1": 4,
        "heading 2": 5,
        "heading 3": 6,
        "heading 4": 7,
        "heading 5": 8,
    }
    return mapping.get(style_name)


def _build_text_style(run) -> dict:
    style: dict = {}
    if run.bold:
        style["bl"] = 1
    if run.italic:
        style["it"] = 1
    if run.underline:
        style["ul"] = {"s": 1}

    font = run.font
    if getattr(font, "strike", False):
        style["st"] = {"s": 1}
    if font is not None and font.name:
        style["ff"] = str(font.name)
    if font is not None and font.size is not None:
        try:
            style["fs"] = float(font.size.pt)
        except Exception:
            pass

    color = getattr(getattr(font, "color", None), "rgb", None)
    if color:
        style["cl"] = {"rgb": f"#{color}"}

    style.update(_build_text_style_from_xml_properties(getattr(getattr(run, "_element", None), "rPr", None)))
    return style


def _twips_to_points(value: str | None) -> float | None:
    if value is None:
        return None
    try:
        return float(value) / 20.0
    except Exception:
        return None


def _map_number_format_to_glyph_type(num_format: str | None) -> int:
    mapping = {
        "bullet": 0,
        "decimal": 2,
        "decimalZero": 3,
        "upperLetter": 4,
        "lowerLetter": 5,
        "upperRoman": 6,
        "lowerRoman": 7,
    }
    return mapping.get(str(num_format or ""), 0)


def _extract_list_level_from_style_id(style_id: str | None) -> int:
    if not style_id:
        return 0
    suffix = ""
    for char in reversed(style_id):
        if char.isdigit():
            suffix = char + suffix
        else:
            break
    if not suffix:
        return 0
    return max(int(suffix) - 1, 0)


def _build_list_catalog(document: Document) -> dict[str, _ListDefinitionInfo]:
    numbering_part = getattr(getattr(document, "part", None), "numbering_part", None)
    if numbering_part is None:
        return {}

    root = etree.fromstring(numbering_part.blob)
    list_catalog: dict[str, _ListDefinitionInfo] = {}
    for abstract_num in root.xpath("./w:abstractNum", namespaces=_WORD_NAMESPACES):
        p_styles = abstract_num.xpath(".//w:pStyle/@w:val", namespaces=_WORD_NAMESPACES)
        if not p_styles:
            continue

        num_format = next(iter(abstract_num.xpath(".//w:numFmt/@w:val", namespaces=_WORD_NAMESPACES)), None)
        glyph_text = next(iter(abstract_num.xpath(".//w:lvlText/@w:val", namespaces=_WORD_NAMESPACES)), None)
        start_number = next(iter(abstract_num.xpath(".//w:start/@w:val", namespaces=_WORD_NAMESPACES)), "1")
        left_indent = _points_to_pixels(
            _twips_to_points(next(iter(abstract_num.xpath(".//w:ind/@w:left", namespaces=_WORD_NAMESPACES)), None))
        )
        hanging_indent = _points_to_pixels(
            _twips_to_points(next(iter(abstract_num.xpath(".//w:ind/@w:hanging", namespaces=_WORD_NAMESPACES)), None))
        )

        for style_id in p_styles:
            level = _extract_list_level_from_style_id(style_id)
            list_type = _BULLET_LIST_TYPE if num_format == "bullet" else _ORDER_LIST_TYPE
            list_catalog[str(style_id)] = _ListDefinitionInfo(
                style_id=str(style_id),
                level=level,
                list_type=list_type,
                glyph_format=str(glyph_text or ("%1." if list_type == _ORDER_LIST_TYPE else "")),
                glyph_symbol=str(glyph_text or "") if list_type == _BULLET_LIST_TYPE else None,
                glyph_type=_map_number_format_to_glyph_type(num_format),
                start_number=int(start_number or 1),
                indent_start=left_indent,
                hanging=hanging_indent,
            )

    return list_catalog


def _build_list_definition(info: _ListDefinitionInfo) -> dict:
    paragraph_properties: dict = {}
    if info.get("indent_start") is not None:
        paragraph_properties["indentStart"] = {"v": float(info["indent_start"])}
    if info.get("hanging") is not None:
        paragraph_properties["hanging"] = {"v": float(info["hanging"])}

    return {
        "listType": info["list_type"],
        "nestingLevel": [
            {
                "bulletAlignment": 1,
                "glyphFormat": info["glyph_format"],
                "startNumber": int(info["start_number"]),
                "glyphType": int(info["glyph_type"]),
                "glyphSymbol": info.get("glyph_symbol"),
                "paragraphProperties": paragraph_properties or None,
            }
        ],
    }


def _apply_list_metadata(
    paragraph,
    paragraph_record: dict,
    lists: dict[str, dict],
    list_catalog: dict[str, _ListDefinitionInfo],
    active_list_ids: dict[str, str],
) -> None:
    style = getattr(paragraph, "style", None)
    style_id = getattr(style, "style_id", None)
    list_info = list_catalog.get(str(style_id or ""))
    if list_info is None:
        return

    level = int(list_info["level"])
    list_key = f"{list_info['style_id']}:{level}"
    list_id = active_list_ids.get(list_key)
    if list_id is None:
        list_id = f"list-{uuid4().hex}"
        active_list_ids[list_key] = list_id
        lists[list_id] = _build_list_definition(list_info)

    paragraph_record["bullet"] = {
        "listType": list_info["list_type"],
        "listId": list_id,
        "nestingLevel": level,
        "textStyle": {},
    }


def _build_paragraph_style(paragraph) -> dict:
    style = _copy_default_paragraph_style()
    named_style_type = _map_named_style_type(paragraph)
    if named_style_type is not None:
        style["namedStyleType"] = named_style_type
    horizontal_align = _map_paragraph_alignment(paragraph)
    if horizontal_align is not None:
        style["horizontalAlign"] = horizontal_align

    paragraph_format = getattr(paragraph, "paragraph_format", None)
    if paragraph_format is None:
        return _apply_paragraph_xml_style(style, paragraph)

    if paragraph_format.space_before is not None:
        try:
            style["spaceAbove"] = {"v": _points_to_pixels(float(paragraph_format.space_before.pt))}
        except Exception:
            pass
    if paragraph_format.space_after is not None:
        try:
            style["spaceBelow"] = {"v": _points_to_pixels(float(paragraph_format.space_after.pt))}
        except Exception:
            pass
    if paragraph_format.line_spacing is not None:
        try:
            line_spacing = float(paragraph_format.line_spacing)
            if line_spacing > 0:
                style["lineSpacing"] = line_spacing
        except Exception:
            pass

    first_line_indent = _to_pixels(paragraph_format.first_line_indent)
    if first_line_indent is not None:
        if first_line_indent >= 0:
            style["indentFirstLine"] = {"v": first_line_indent}
        else:
            style["hanging"] = {"v": abs(first_line_indent)}

    left_indent = _to_pixels(paragraph_format.left_indent)
    if left_indent is not None:
        style["indentStart"] = {"v": left_indent}

    right_indent = _to_pixels(paragraph_format.right_indent)
    if right_indent is not None:
        style["indentEnd"] = {"v": right_indent}

    if paragraph_format.keep_with_next is not None:
        style["keepNext"] = 1 if paragraph_format.keep_with_next else 0
    if paragraph_format.keep_together is not None:
        style["keepLines"] = 1 if paragraph_format.keep_together else 0
    if paragraph_format.widow_control is not None:
        style["widowControl"] = 1 if paragraph_format.widow_control else 0

    return _apply_paragraph_xml_style(style, paragraph)


def _build_paragraph_border(border_element) -> dict | None:
    border = _build_table_cell_border(border_element)
    if border is None:
        return None
    space_raw = border_element.get(f"{{{_WORD_NAMESPACE}}}space")
    try:
        border["padding"] = int(space_raw) if space_raw is not None else 0
    except Exception:
        border["padding"] = 0
    return border


def _parse_on_off_value(value: str | None, default: bool = True) -> int:
    normalized = str(value or "").strip().lower()
    if not normalized:
        return 1 if default else 0
    if normalized in {"0", "false", "off", "no"}:
        return 0
    return 1


def _map_word_highlight_color(value: str | None) -> dict | None:
    color = _WORD_HIGHLIGHT_COLORS.get(str(value or ""))
    if color is None:
        return None
    return {"rgb": color}


def _build_text_style_from_xml_properties(run_properties) -> dict:
    if run_properties is None:
        return {}

    style: dict = {}
    if next(iter(run_properties.xpath('./*[local-name()="b"]')), None) is not None:
        style["bl"] = _parse_on_off_value(
            next(iter(run_properties.xpath('./*[local-name()="b"]')), None).get(f"{{{_WORD_NAMESPACE}}}val"),
            default=True,
        )
    if next(iter(run_properties.xpath('./*[local-name()="i"]')), None) is not None:
        style["it"] = _parse_on_off_value(
            next(iter(run_properties.xpath('./*[local-name()="i"]')), None).get(f"{{{_WORD_NAMESPACE}}}val"),
            default=True,
        )
    if next(iter(run_properties.xpath('./*[local-name()="u"]')), None) is not None:
        style["ul"] = {"s": _parse_on_off_value(
            next(iter(run_properties.xpath('./*[local-name()="u"]')), None).get(f"{{{_WORD_NAMESPACE}}}val"),
            default=True,
        )}
    if next(iter(run_properties.xpath('./*[local-name()="strike"]')), None) is not None:
        style["st"] = {"s": _parse_on_off_value(
            next(iter(run_properties.xpath('./*[local-name()="strike"]')), None).get(f"{{{_WORD_NAMESPACE}}}val"),
            default=True,
        )}

    font_element = next(iter(run_properties.xpath('./*[local-name()="rFonts"]')), None)
    if font_element is not None:
        font_name = (
            font_element.get(f"{{{_WORD_NAMESPACE}}}ascii")
            or font_element.get(f"{{{_WORD_NAMESPACE}}}hAnsi")
            or font_element.get(f"{{{_WORD_NAMESPACE}}}cs")
        )
        if font_name:
            style["ff"] = str(font_name)

    size_element = next(iter(run_properties.xpath('./*[local-name()="sz"]')), None)
    if size_element is not None:
        try:
            style["fs"] = float(size_element.get(f"{{{_WORD_NAMESPACE}}}val")) / 2.0
        except Exception:
            pass

    color_element = next(iter(run_properties.xpath('./*[local-name()="color"]')), None)
    color = _parse_word_color(color_element.get(f"{{{_WORD_NAMESPACE}}}val") if color_element is not None else None)
    if color is not None:
        style["cl"] = color

    highlight_element = next(iter(run_properties.xpath('./*[local-name()="highlight"]')), None)
    highlight = _map_word_highlight_color(
        highlight_element.get(f"{{{_WORD_NAMESPACE}}}val") if highlight_element is not None else None
    )
    if highlight is not None:
        style["bg"] = highlight

    vertical_align_element = next(iter(run_properties.xpath('./*[local-name()="vertAlign"]')), None)
    vertical_align = vertical_align_element.get(f"{{{_WORD_NAMESPACE}}}val") if vertical_align_element is not None else None
    if vertical_align == "subscript":
        style["va"] = 2
    elif vertical_align == "superscript":
        style["va"] = 3

    return style


def _map_tab_stop_alignment(value: str | None) -> int:
    mapping = {
        "left": 1,
        "start": 1,
        "center": 2,
        "right": 3,
        "end": 3,
        "decimal": 3,
    }
    return mapping.get(str(value or ""), 1)


def _build_paragraph_tab_stops(paragraph_properties) -> list[dict] | None:
    tabs_element = next(iter(paragraph_properties.xpath('./*[local-name()="tabs"]')), None)
    if tabs_element is None:
        return None

    tab_stops: list[dict] = []
    for tab_element in tabs_element.xpath('./*[local-name()="tab"]'):
        offset = _twips_to_points(tab_element.get(f"{{{_WORD_NAMESPACE}}}pos"))
        if offset is None:
            continue
        tab_stops.append({
            "offset": _points_to_pixels(offset),
            "alignment": _map_tab_stop_alignment(tab_element.get(f"{{{_WORD_NAMESPACE}}}val")),
        })
    return tab_stops or None


def _apply_paragraph_xml_style(style: dict, paragraph) -> dict:
    paragraph_element = getattr(paragraph, "_p", None)
    paragraph_properties = getattr(paragraph_element, "pPr", None)
    if paragraph_properties is None:
        return style

    shading = next(iter(paragraph_properties.xpath('./*[local-name()="shd"]')), None)
    fill = shading.get(f"{{{_WORD_NAMESPACE}}}fill") if shading is not None else None
    background = _parse_word_color(fill)
    if background is not None:
        style["shading"] = background

    borders = next(iter(paragraph_properties.xpath('./*[local-name()="pBdr"]')), None)
    if borders is not None:
        for xml_name, field_name in (
            ("top", "borderTop"),
            ("left", "borderLeft"),
            ("bottom", "borderBottom"),
            ("right", "borderRight"),
            ("between", "borderBetween"),
        ):
            border_element = next(iter(borders.xpath(f'./*[local-name()="{xml_name}"]')), None)
            border = _build_paragraph_border(border_element)
            if border is not None:
                style[field_name] = border

    tab_stops = _build_paragraph_tab_stops(paragraph_properties)
    if tab_stops is not None:
        style["tabStops"] = tab_stops

    text_style = _build_text_style_from_xml_properties(
        next(iter(paragraph_properties.xpath('./*[local-name()="rPr"]')), None)
    )
    if text_style:
        style["textStyle"] = text_style

    spacing_element = next(iter(paragraph_properties.xpath('./*[local-name()="spacing"]')), None)
    if spacing_element is not None:
        line_value = _points_to_pixels(_twips_to_points(spacing_element.get(f"{{{_WORD_NAMESPACE}}}line")))
        if line_value is not None:
            style["lineSpacing"] = line_value
        line_rule = str(spacing_element.get(f"{{{_WORD_NAMESPACE}}}lineRule") or "").strip().lower()
        if line_rule == "atleast":
            style["spacingRule"] = 1
        elif line_rule == "exact":
            style["spacingRule"] = 2
        elif line_rule:
            style["spacingRule"] = 0

    if next(iter(paragraph_properties.xpath('./*[local-name()="bidi"]')), None) is not None:
        style["direction"] = 2

    if next(iter(paragraph_properties.xpath('./*[local-name()="suppressAutoHyphens"]')), None) is not None:
        style["suppressHyphenation"] = _parse_on_off_value(
            next(iter(paragraph_properties.xpath('./*[local-name()="suppressAutoHyphens"]')), None).get(
                f"{{{_WORD_NAMESPACE}}}val"
            ),
            default=True,
        )

    return style


def _append_text_with_paragraph_breaks(
    text: str,
    data_stream_parts: list[str],
    paragraph_records: list[dict],
    paragraph_style: dict,
    cursor: int,
    text_style: dict | None = None,
    text_runs: list[dict] | None = None,
) -> int:
    if not text:
        return cursor

    start = cursor
    data_stream_parts.append(text)
    cursor += len(text)
    if text_style and text_runs is not None:
        text_runs.append({
            "st": start,
            "ed": cursor,
            "ts": dict(text_style),
        })
    return cursor


def _extract_note_entries_from_xml(xml_bytes: bytes, entry_name: str) -> dict[str, str]:
    root = etree.fromstring(xml_bytes)
    notes: dict[str, str] = {}
    for note in root.xpath(f'./w:{entry_name}', namespaces=_WORD_NAMESPACES):
        note_id = str(note.get(f"{{{_WORD_NAMESPACE}}}id") or "")
        note_type = note.get(f"{{{_WORD_NAMESPACE}}}type")
        if not note_id or note_type in {"separator", "continuationSeparator"}:
            continue

        fragments: list[str] = []
        for paragraph in note.xpath('./w:p', namespaces=_WORD_NAMESPACES):
            text = "".join(paragraph.xpath('.//w:t/text()', namespaces=_WORD_NAMESPACES)).strip()
            if text:
                fragments.append(text)
        if fragments:
            notes[note_id] = "\n".join(fragments)
    return notes


def _get_html_preview_dir(user_id: int, conversation_id: str, folder_name: str, cache_key: str) -> Path:
    return get_conversation_dir(user_id, conversation_id) / ".agent" / "preview_cache" / folder_name / cache_key


def _build_html_preview_relative_path(folder_name: str, cache_key: str) -> str:
    return f".agent/preview_cache/{folder_name}/{cache_key}/index.html"


def _build_preview_cache_key(source_bytes: bytes) -> str:
    return hashlib.sha256(source_bytes).hexdigest()


def _get_doc_html_preview_dir(user_id: int, conversation_id: str, cache_key: str) -> Path:
    return _get_html_preview_dir(user_id, conversation_id, "docx-parser-html", cache_key)


def _get_presentation_html_preview_dir(user_id: int, conversation_id: str, cache_key: str) -> Path:
    return _get_html_preview_dir(user_id, conversation_id, "ppt-parser-html", cache_key)


def _get_sheet_html_preview_dir(user_id: int, conversation_id: str, cache_key: str) -> Path:
    return _get_html_preview_dir(user_id, conversation_id, "xlsx-parser-html", cache_key)


def _get_soffice_preview_env() -> dict[str, str]:
    try:
        return get_soffice_env()
    except AttributeError:
        env = os.environ.copy()
        env["SAL_USE_VCLPLUGIN"] = "svp"
        return env


def _convert_office_to_html_preview(
    user_id: int,
    conversation_id: str,
    resolved: Path,
    source_bytes: bytes,
    *,
    preview_dir: Path,
    output_relative_path: str,
    missing_dependency_message: str,
    failure_prefix: str,
) -> tuple[str, bool]:
    if (preview_dir / "index.html").exists():
        return output_relative_path, True

    preview_dir.parent.mkdir(parents=True, exist_ok=True)
    temp_dir = Path(tempfile.mkdtemp(prefix=f"{preview_dir.name}-", dir=preview_dir.parent))
    source_copy = temp_dir / resolved.name
    source_copy.write_bytes(source_bytes)

    try:
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "html",
                "--outdir",
                str(temp_dir),
                str(source_copy),
            ],
            check=True,
            capture_output=True,
            text=True,
            env=_get_soffice_preview_env(),
        )

        generated_entries = sorted(
            candidate for candidate in temp_dir.iterdir()
            if candidate.is_file() and candidate.suffix.lower() in {".html", ".htm"}
        )
        if not generated_entries:
            raise ValueError("LibreOffice did not produce an HTML preview file")

        generated_entry = generated_entries[0]
        target_entry = temp_dir / "index.html"
        if generated_entry != target_entry:
            generated_entry.rename(target_entry)
        source_copy.unlink(missing_ok=True)

        if preview_dir.exists():
            shutil.rmtree(preview_dir)
        shutil.move(str(temp_dir), str(preview_dir))
    except FileNotFoundError as exc:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise ValueError(missing_dependency_message) from exc
    except subprocess.CalledProcessError as exc:
        shutil.rmtree(temp_dir, ignore_errors=True)
        stderr = (exc.stderr or exc.stdout or "").strip()
        detail = f": {stderr}" if stderr else ""
        raise ValueError(f"{failure_prefix}{detail}") from exc
    except Exception:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise

    return output_relative_path, False


def _convert_docx_to_html_preview(
    user_id: int,
    conversation_id: str,
    file_path: str,
    resolved: Path,
    source_bytes: bytes,
) -> tuple[str, bool]:
    cache_key = _build_preview_cache_key(source_bytes)
    return _convert_office_to_html_preview(
        user_id,
        conversation_id,
        resolved,
        source_bytes,
        preview_dir=_get_doc_html_preview_dir(user_id, conversation_id, cache_key),
        output_relative_path=_build_html_preview_relative_path("docx-parser-html", cache_key),
        missing_dependency_message="LibreOffice soffice is required to preview document files",
        failure_prefix="Failed to generate document HTML preview",
    )


def _convert_sheet_to_html_preview(
    user_id: int,
    conversation_id: str,
    file_path: str,
    resolved: Path,
    source_bytes: bytes,
) -> tuple[str, bool]:
    cache_key = _build_preview_cache_key(source_bytes)
    preview_dir = _get_sheet_html_preview_dir(user_id, conversation_id, cache_key)
    result = _convert_office_to_html_preview(
        user_id,
        conversation_id,
        resolved,
        source_bytes,
        preview_dir=preview_dir,
        output_relative_path=_build_html_preview_relative_path("xlsx-parser-html", cache_key),
        missing_dependency_message="LibreOffice soffice is required to preview spreadsheet files with charts",
        failure_prefix="Failed to generate spreadsheet HTML preview",
    )
    _polish_sheet_html_preview(preview_dir / "index.html")
    return result


def _polish_sheet_html_preview(preview_html: Path) -> None:
    html_text = preview_html.read_text(encoding="utf-8")
    marker = "home-harness-sheet-html-preview"
    if marker in html_text:
        return

    style = f"""<style id="{marker}">
      :root {{
        color-scheme: light;
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Arial, sans-serif;
        background: #ffffff;
        color: #111827;
      }}
      html,
      body {{
        min-height: 100%;
        margin: 0;
        background: #ffffff;
      }}
      body {{
        box-sizing: border-box;
        padding: 18px 22px 32px;
        overflow: auto;
        font-size: 13px;
        line-height: 1.4;
      }}
      body > hr:first-of-type,
      body > center:first-of-type,
      body > center:first-of-type + hr {{
        display: none !important;
      }}
      h1,
      h2,
      h3 {{
        color: #111827;
        font-style: normal;
        letter-spacing: 0;
      }}
      h1 {{
        margin: 0 0 14px;
        font-size: 20px;
        line-height: 1.3;
        font-weight: 700;
      }}
      table {{
        width: auto !important;
        max-width: 100%;
        margin: 0 0 20px;
        border-collapse: collapse !important;
        border-spacing: 0 !important;
        background: #ffffff;
      }}
      th,
      td {{
        min-width: 72px;
        height: 24px;
        padding: 3px 8px !important;
        border: 1px solid #d8dee8 !important;
        box-sizing: border-box;
        color: #111827;
        font-size: 13px;
        line-height: 1.35;
        vertical-align: middle;
        white-space: pre-wrap;
        overflow-wrap: anywhere;
      }}
      th {{
        background: #f4f7fb;
        font-weight: 600;
      }}
      img,
      svg {{
        display: block;
        max-width: min(100%, 1120px);
        height: auto;
        margin: 18px 0 24px;
        border: 1px solid #d8dee8;
        background: #ffffff;
      }}
      a {{
        color: #2563eb;
      }}
    </style>"""

    lower_html = html_text.lower()
    head_end = lower_html.find("</head>")
    if head_end != -1:
        html_text = f"{html_text[:head_end]}{style}\n{html_text[head_end:]}"
    else:
        html_start = lower_html.find("<html")
        html_open_end = lower_html.find(">", html_start) if html_start != -1 else -1
        if html_open_end != -1:
            html_text = f"{html_text[:html_open_end + 1]}<head>{style}</head>\n{html_text[html_open_end + 1:]}"
        else:
            html_text = f"{style}\n{html_text}"

    preview_html.write_text(html_text, encoding="utf-8")


def _prefer_sheet_html_preview(
    capabilities: dict | None,
) -> bool:
    detected = capabilities.get("detected_unsupported_features") if isinstance(capabilities, dict) else None
    if not isinstance(detected, list):
        return False
    return any(str(feature) in _XLSX_HTML_PREVIEW_FEATURES for feature in detected)


def _replace_sheet_drawing_warning(warnings: list[str], replacement: str) -> list[str]:
    drawing_warning = _XLSX_UNSUPPORTED_FEATURE_WARNINGS["charts"]
    next_warnings = [warning for warning in warnings if warning != drawing_warning]
    if replacement not in next_warnings:
        next_warnings.append(replacement)
    return next_warnings


def _convert_presentation_to_html_preview(
    user_id: int,
    conversation_id: str,
    file_path: str,
    resolved: Path,
    source_bytes: bytes,
) -> tuple[str, bool]:
    cache_key = _build_preview_cache_key(source_bytes)
    preview_dir = _get_presentation_html_preview_dir(user_id, conversation_id, cache_key)
    output_relative_path = _build_html_preview_relative_path("ppt-parser-html", cache_key)
    if (preview_dir / "index.html").exists():
        return output_relative_path, True

    preview_dir.parent.mkdir(parents=True, exist_ok=True)
    temp_dir = Path(tempfile.mkdtemp(prefix=f"{preview_dir.name}-", dir=preview_dir.parent))
    source_copy = temp_dir / resolved.name
    source_copy.write_bytes(source_bytes)

    try:
        generated_svgs = _export_presentation_slides_to_svg(source_copy, temp_dir)
        if not generated_svgs:
            raise ValueError("LibreOffice did not produce slide SVG previews")

        (temp_dir / "index.html").write_text(
            _build_presentation_svg_preview_html(resolved.name, generated_svgs),
            encoding="utf-8",
        )
        source_copy.unlink(missing_ok=True)

        if preview_dir.exists():
            shutil.rmtree(preview_dir)
        shutil.move(str(temp_dir), str(preview_dir))
    except FileNotFoundError as exc:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise ValueError("LibreOffice soffice is required to preview presentation files") from exc
    except subprocess.CalledProcessError as exc:
        shutil.rmtree(temp_dir, ignore_errors=True)
        stderr = (exc.stderr or exc.stdout or "").strip()
        detail = f": {stderr}" if stderr else ""
        raise ValueError(f"Failed to generate presentation SVG preview{detail}") from exc
    except Exception:
        shutil.rmtree(temp_dir, ignore_errors=True)
        raise

    return output_relative_path, False


def _build_presentation_svg_preview_html(file_name: str, svg_paths: list[Path]) -> str:
    total_slides = len(svg_paths)
    slides_markup = "\n".join(
        (
            f'<section class="slide-sheet" aria-label="PPT slide {index}">'
            f'<div class="slide-toolbar">'
            f'<div class="slide-counter">{index}/{total_slides}</div>'
            f'<button type="button" class="slide-regenerate-button" data-slide-index="{index}">'
            f"重新生成第{index}页"
            f"</button>"
            f"</div>"
            f'<div class="slide-canvas">'
            f"{_read_presentation_svg_markup(path)}"
            f"</div>"
            f"</section>"
        )
        for index, path in enumerate(svg_paths, start=1)
    )
    return f"""<!DOCTYPE html>
<html lang="en">
  <head>
    <meta charset="utf-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1" />
    <title>PPT Preview</title>
    <style>
      :root {{
        color-scheme: light;
        font-family: "Segoe UI", Arial, sans-serif;
        background: #f5f5f4;
        color: #18181b;
      }}
      * {{
        box-sizing: border-box;
      }}
      body {{
        margin: 0;
        padding: 32px 40px 48px;
        background:
          radial-gradient(circle at top, rgba(255,255,255,0.92), rgba(245,245,244,0.98)),
          #f5f5f4;
      }}
      .deck {{
        max-width: 1200px;
        margin: 0 auto;
        display: flex;
        flex-direction: column;
        gap: 28px;
      }}
      .slide-sheet {{
        display: flex;
        flex-direction: column;
        gap: 18px;
      }}
      .slide-toolbar {{
        display: flex;
        align-items: center;
        justify-content: space-between;
        gap: 16px;
      }}
      .slide-counter {{
        font-size: 18px;
        font-weight: 500;
        color: #334155;
        letter-spacing: -0.01em;
      }}
      .slide-regenerate-button {{
        border: 1px solid rgba(15, 23, 42, 0.08);
        border-radius: 18px;
        background: rgba(255, 255, 255, 0.92);
        color: #1e293b;
        padding: 12px 18px;
        font-size: 14px;
        line-height: 1;
        cursor: pointer;
        transition: background 120ms ease, border-color 120ms ease, color 120ms ease, opacity 120ms ease;
      }}
      .slide-regenerate-button:hover {{
        background: #ffffff;
        border-color: rgba(15, 23, 42, 0.14);
      }}
      .slide-regenerate-button:disabled {{
        cursor: not-allowed;
        opacity: 0.48;
      }}
      .slide-canvas {{
        width: 100%;
      }}
      .slide-canvas > svg {{
        display: block;
        width: 100%;
        height: auto;
        background: white;
        border-radius: 24px;
        box-shadow: 0 16px 40px rgba(15, 23, 42, 0.08);
      }}
    </style>
  </head>
  <body>
    <main class="deck">
      {slides_markup}
    </main>
    <script>
      (() => {{
        const messageType = "home_harness_ppt_regenerate_slide";
        const stateMessageType = "home_harness_ppt_regenerate_state";
        const buttons = Array.from(document.querySelectorAll("[data-slide-index]"));
        const setDisabled = (disabled) => {{
          buttons.forEach((button) => {{
            button.disabled = !!disabled;
          }});
        }};

        buttons.forEach((button) => {{
          button.addEventListener("click", () => {{
            if (button.disabled) {{
              return;
            }}
            const slideIndex = Number(button.getAttribute("data-slide-index") || "0");
            if (!Number.isInteger(slideIndex) || slideIndex <= 0) {{
              return;
            }}
            window.parent.postMessage({{ type: messageType, slideIndex }}, "*");
          }});
        }});

        window.addEventListener("message", (event) => {{
          if (!event.data || event.data.type !== stateMessageType) {{
            return;
          }}
          setDisabled(Boolean(event.data.disabled));
        }});
      }})();
    </script>
  </body>
</html>
"""


def _read_presentation_svg_markup(svg_path: Path) -> str:
    return svg_path.read_text(encoding="utf-8")


def _get_presentation_slide_count(resolved: Path) -> int:
    try:
        with ZipFile(resolved, "r") as archive:
            slide_names = [
                name for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ]
    except Exception as exc:
        raise ValueError("Failed to inspect presentation slides for preview") from exc

    return len(slide_names)


def _sanitize_presentation_svg_for_static_preview(svg_path: Path, slide_index: int) -> None:
    parser = etree.XMLParser(remove_blank_text=False, recover=True)
    root = etree.parse(str(svg_path), parser).getroot()
    modified = False

    for script in root.xpath(".//*[local-name()='script']"):
        parent = script.getparent()
        if parent is not None:
            parent.remove(script)
            modified = True

    slide_group_nodes = root.xpath(".//*[local-name()='g' and @class='SlideGroup']")
    if not slide_group_nodes:
        if modified:
            svg_path.write_bytes(
                etree.tostring(root, xml_declaration=True, encoding="UTF-8")
            )
        return

    slide_group = slide_group_nodes[0]
    target_container_id = f"container-id{slide_index}"
    target_wrappers = slide_group.xpath(f"./*[.//*[@id='{target_container_id}']]")

    if target_wrappers:
        target_wrapper = target_wrappers[0]
        for child in list(slide_group):
            if child is not target_wrapper:
                slide_group.remove(child)
                modified = True
        if target_wrapper.attrib.get("visibility") == "hidden":
            target_wrapper.attrib["visibility"] = "visible"
            modified = True
        page_nodes = target_wrapper.xpath(".//*[local-name()='g' and @class='Page']")
        if page_nodes:
            page_node = page_nodes[0]
            background_defs_nodes = page_node.xpath("./*[local-name()='defs' and @class='SlideBackground']")
            if background_defs_nodes:
                background_defs = background_defs_nodes[0]
                visible_background_nodes = [
                    deepcopy(child)
                    for child in background_defs
                    if isinstance(child.tag, str)
                ]
                if visible_background_nodes:
                    insert_at = page_node.index(background_defs)
                    for offset, child in enumerate(visible_background_nodes):
                        page_node.insert(insert_at + offset, child)
                    modified = True

    if modified:
        svg_path.write_bytes(
            etree.tostring(root, xml_declaration=True, encoding="UTF-8")
        )


def _export_presentation_slides_to_svg(resolved: Path, temp_dir: Path) -> list[Path]:
    slide_count = _get_presentation_slide_count(resolved)
    if slide_count <= 0:
        raise ValueError("Presentation contains no slides to preview")

    generated_svgs: list[Path] = []
    env = _get_soffice_preview_env()

    for slide_index in range(1, slide_count + 1):
        generated_entry = temp_dir / f"{resolved.stem}.svg"
        generated_entry.unlink(missing_ok=True)
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                (
                    'svg:impress_svg_Export:'
                    f'{{"PageNumber":{{"type":"long","value":"{slide_index}"}}}}'
                ),
                "--outdir",
                str(temp_dir),
                str(resolved),
            ],
            check=True,
            capture_output=True,
            text=True,
            env=env,
        )
        if not generated_entry.exists():
            raise ValueError(f"LibreOffice did not produce SVG for slide {slide_index}")
        _sanitize_presentation_svg_for_static_preview(generated_entry, slide_index)
        slide_target = temp_dir / f"slide-{slide_index}.svg"
        generated_entry.replace(slide_target)
        generated_svgs.append(slide_target)

    return generated_svgs


def _load_sheet_snapshot(resolved: Path, extension: str) -> dict:
    if extension == "xls":
        return {
            "kind": "sheet",
            "sheets": [],
            "warning": "Legacy .xls files are preview-only in this version.",
        }

    if extension == "csv":
        with resolved.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = [list(row) for row in csv.reader(handle)]
        return {
            "kind": "sheet",
            "sheets": [_build_sparse_sheet_snapshot(resolved.stem or "Sheet1", _trim_sheet_cells(rows))],
            "styles": {},
        }

    if load_workbook is None:
        raise ValueError("openpyxl is required to open spreadsheet files")

    workbook = load_workbook(resolved, data_only=False)
    sheets: list[dict] = []
    for worksheet in workbook.worksheets:
        sheets.append(_build_workbook_sheet_snapshot(worksheet))

    return {
        "kind": "sheet",
        "sheets": sheets or [_build_sparse_sheet_snapshot("Sheet1", [[]])],
        "styles": {},
    }


def _load_presentation_snapshot() -> dict:
    return {
        "kind": "presentation",
        "slides": [],
    }


def _load_snapshot(file_kind: OfficeFileKind, resolved: Path, extension: str) -> dict:
    if file_kind == "sheet":
        return _load_sheet_snapshot(resolved, extension)
    return _load_presentation_snapshot()


def _get_office_source_mime_type(file_kind: OfficeFileKind, extension: str) -> str | None:
    if file_kind == "sheet":
        if extension == "xlsx":
            return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if extension == "xls":
            return "application/vnd.ms-excel"
        if extension == "csv":
            return "text/csv"
    return None


def _archive_contains_rich_text_runs(resolved: Path) -> bool:
    try:
        with ZipFile(resolved, "r") as archive:
            candidate_names = [
                name for name in archive.namelist()
                if name == "xl/sharedStrings.xml" or (name.startswith("xl/worksheets/") and name.endswith(".xml"))
            ]
            for name in candidate_names:
                try:
                    root = etree.fromstring(archive.read(name))
                except Exception:
                    continue
                if root.xpath('//*[local-name()="si"]/*[local-name()="r"]'):
                    return True
                if root.xpath('//*[local-name()="is"]/*[local-name()="r"]'):
                    return True
    except Exception:
        return False
    return False


def _scan_xlsx_unsupported_features(resolved: Path) -> list[str]:
    if load_workbook is None:
        raise ValueError("openpyxl is required to open spreadsheet files")

    workbook = load_workbook(resolved, data_only=False)
    detected: set[str] = set()

    for worksheet in workbook.worksheets:
        if any(cell.comment is not None for row in worksheet.iter_rows() for cell in row):
            detected.add("comments")
        if len(worksheet.conditional_formatting) > 0:
            detected.add("conditional_formatting")
        if getattr(getattr(worksheet, "data_validations", None), "dataValidation", None):
            if len(worksheet.data_validations.dataValidation) > 0:
                detected.add("data_validation")
        if getattr(worksheet, "_charts", None):
            if len(worksheet._charts) > 0:
                detected.add("charts")
        if getattr(worksheet, "_images", None):
            if len(worksheet._images) > 0:
                detected.add("images")

    try:
        with ZipFile(resolved, "r") as archive:
            archive_names = archive.namelist()
            if any("vbaProject" in name or name.lower().endswith(".bin") and "vba" in name.lower() for name in archive_names):
                detected.add("macros")

            drawing_rel_names = [
                name for name in archive_names
                if name.startswith("xl/drawings/_rels/") and name.endswith(".rels")
            ]
            for name in drawing_rel_names:
                try:
                    root = etree.fromstring(archive.read(name))
                except Exception:
                    continue
                for relationship in root.xpath('//*[local-name()="Relationship"]'):
                    target = str(relationship.get("Target") or "").lower()
                    if "/charts/" in target:
                        detected.add("charts")
                    elif "/media/" in target:
                        detected.add("images")
                    elif target:
                        detected.add("shapes")
            if (
                "shapes" not in detected
                and any(name.startswith("xl/drawings/") and name.endswith(".xml") for name in archive_names)
                and "charts" not in detected
                and "images" not in detected
            ):
                detected.add("shapes")
    except Exception:
        pass

    if _archive_contains_rich_text_runs(resolved):
        detected.add("rich_text")

    return [feature for feature in _XLSX_UNSUPPORTED_FEATURE_ORDER if feature in detected]


def _build_xlsx_open_metadata(resolved: Path) -> tuple[list[str], dict]:
    detected_unsupported_features = _scan_xlsx_unsupported_features(resolved)
    warnings = ["预览中会保留公式，但不会重新计算结果。"]
    seen_warnings = set(warnings)
    for feature in detected_unsupported_features:
        warning = _XLSX_UNSUPPORTED_FEATURE_WARNINGS[feature]
        if warning not in seen_warnings:
            warnings.append(warning)
            seen_warnings.add(warning)

    capabilities = {
        "can_edit": False,
        "can_save": False,
        "supports_styles": True,
        "supports_merges": True,
        "supports_freeze": True,
        "supports_comments": False,
        "supports_conditional_formatting": False,
        "supports_data_validation": False,
        "supports_charts": False,
        "supports_images": False,
        "supports_shapes": False,
        "supports_macros": False,
        "supports_rich_text": False,
        "detected_unsupported_features": detected_unsupported_features,
    }
    return warnings, capabilities


def _get_sheet_open_metadata(resolved: Path, extension: str) -> tuple[list[str], dict | None]:
    if extension == "xlsx":
        return _build_xlsx_open_metadata(resolved)
    if extension == "csv":
        return (
            ["CSV 会以单个工作表打开，且不包含工作簿级能力。"],
            {
                "can_edit": False,
                "can_save": False,
                "supports_styles": False,
                "supports_merges": False,
                "supports_freeze": False,
            },
        )
    if extension == "xls":
        return (
            ["Legacy .xls files are preview-only in this version."],
            {
                "can_edit": False,
                "can_save": False,
                "supports_styles": False,
                "supports_merges": False,
                "supports_freeze": False,
            },
        )
    return ([], None)


def open_workspace_office_session(
    user_id: int,
    conversation_id: str,
    file_path: str,
    *,
    resolved_path: Path | None = None,
) -> dict:
    started_at = time.perf_counter()
    resolved = resolved_path or get_workspace_file_path(user_id, conversation_id, file_path)
    if resolved is None:
        raise FileNotFoundError(file_path)

    cleanup_expired_workspace_office_sessions(user_id, conversation_id)

    file_kind = _get_file_kind(file_path)
    extension = _get_file_extension(file_path)
    max_source_bytes = max(1, int(getattr(settings, "OFFICE_SESSION_MAX_SOURCE_BYTES", 100 * 1024 * 1024) or 0))
    source_size = resolved.stat().st_size
    if source_size > max_source_bytes:
        raise ValueError(f"Office file exceeds maximum session size of {max_source_bytes} bytes")
    source_bytes = resolved.read_bytes()
    session_id = f"office-{uuid4().hex}"
    engine = "html" if (
        (file_kind == "doc" and extension == "docx")
        or file_kind == "presentation"
    ) else "univer"
    preview_file_path = None
    preview_cache_hit = False
    readonly = True if file_kind == "sheet" else _is_readonly_extension(extension, file_kind)
    unit_id = f"unit-{uuid4().hex}"
    snapshot = None
    conversion_ms = 0.0
    snapshot_ms = 0.0
    metadata_ms = 0.0
    if engine == "html":
        conversion_started_at = time.perf_counter()
        if file_kind == "presentation":
            preview_file_path, preview_cache_hit = _convert_presentation_to_html_preview(
                user_id,
                conversation_id,
                file_path,
                resolved,
                source_bytes,
            )
        else:
            preview_file_path, preview_cache_hit = _convert_docx_to_html_preview(
                user_id,
                conversation_id,
                file_path,
                resolved,
                source_bytes,
            )
        conversion_ms = round((time.perf_counter() - conversion_started_at) * 1000, 2)
        readonly = True
        unit_id = None
    else:
        snapshot_started_at = time.perf_counter()
        snapshot = _load_snapshot(file_kind, resolved, extension)
        snapshot_ms = round((time.perf_counter() - snapshot_started_at) * 1000, 2)

    opened_at = datetime.now(timezone.utc).isoformat()
    record = _build_office_session_record(
        session_id=session_id,
        user_id=user_id,
        conversation_id=conversation_id,
        file_path=file_path,
        file_kind=file_kind,
        file_extension=extension,
        engine=engine,
        readonly=readonly,
        opened_at=opened_at,
    )
    _OFFICE_SESSIONS[session_id] = record
    _persist_office_session_record(user_id, conversation_id, session_id, record)
    response: dict = {
        "session_id": session_id,
        "file_path": file_path,
        "file_kind": file_kind,
        "engine": engine,
        "unit_id": unit_id,
        "readonly": readonly,
    }
    if snapshot is not None:
        response["snapshot"] = snapshot
    if preview_file_path is not None:
        response["preview_file_path"] = preview_file_path

    source_mime_type = _get_office_source_mime_type(file_kind, extension)
    if source_mime_type is not None:
        response["source_blob"] = base64.b64encode(source_bytes).decode("ascii")
        response["source_mime_type"] = source_mime_type
        response["source_encoding"] = "base64"

    if file_kind == "sheet":
        metadata_started_at = time.perf_counter()
        warnings, capabilities = _get_sheet_open_metadata(resolved, extension)
        metadata_ms = round((time.perf_counter() - metadata_started_at) * 1000, 2)
        if extension == "xlsx" and _prefer_sheet_html_preview(capabilities):
            conversion_started_at = time.perf_counter()
            try:
                preview_file_path, preview_cache_hit = _convert_sheet_to_html_preview(
                    user_id,
                    conversation_id,
                    file_path,
                    resolved,
                    source_bytes,
                )
                conversion_ms = round((time.perf_counter() - conversion_started_at) * 1000, 2)
                record["preview_file_path"] = preview_file_path
                response["preview_file_path"] = preview_file_path
                warnings = _replace_sheet_drawing_warning(warnings, _XLSX_HTML_PREVIEW_WARNING)
            except ValueError:
                warnings = _replace_sheet_drawing_warning(warnings, _XLSX_HTML_PREVIEW_FALLBACK_WARNING)
        response["warnings"] = warnings
        response["capabilities"] = capabilities

    elapsed_ms = round((time.perf_counter() - started_at) * 1000, 2)
    logger.info(
        "Opened office session user_id=%s conversation_id=%s file_path=%s file_kind=%s extension=%s engine=%s "
        "elapsed_ms=%s conversion_ms=%s snapshot_ms=%s metadata_ms=%s preview_cache_hit=%s source_bytes=%s",
        user_id,
        conversation_id,
        file_path,
        file_kind,
        extension,
        engine,
        elapsed_ms,
        conversion_ms,
        snapshot_ms,
        metadata_ms,
        preview_cache_hit,
        len(source_bytes),
    )

    return response


def close_workspace_office_session(user_id: int, conversation_id: str, session_id: str) -> dict:
    record = _OFFICE_SESSIONS.get(session_id)
    if record is None:
        record = _load_office_session_record(user_id, conversation_id, session_id)
    if record is None or record["user_id"] != user_id or record["conversation_id"] != conversation_id:
        raise ValueError("Office session not found")

    _OFFICE_SESSIONS.pop(session_id, None)
    _delete_office_session_record(user_id, conversation_id, session_id)
    return {
        "session_id": session_id,
        "closed": True,
    }
