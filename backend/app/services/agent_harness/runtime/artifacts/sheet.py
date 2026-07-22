from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Any

from app.services.agent_harness.runtime.artifacts.base import ArtifactFamily, StructuredArtifactDraft, StructuredSheet

if TYPE_CHECKING:
    from app.services.agent_harness.core.context import HarnessContext


@dataclass(slots=True)
class RenderedSheetArtifact:
    path: Path
    relative_path: str
    metadata: dict[str, Any] = field(default_factory=dict)


class SheetRuntime:
    """Deterministic renderer for structured sheet drafts."""

    def __init__(self, ctx: "HarnessContext") -> None:
        self.ctx = ctx

    def render_draft(self, draft: StructuredArtifactDraft, *, output_root: str | None = None) -> RenderedSheetArtifact:
        from openpyxl import Workbook

        if draft.family != ArtifactFamily.SHEET:
            raise ValueError("SheetRuntime only renders sheet drafts")
        sheets = [_coerce_sheet(item) for item in draft.sheets]
        if not sheets:
            raise ValueError("structured sheet draft must contain at least one sheet")
        workbook = Workbook()
        default = workbook.active
        workbook.remove(default)
        for sheet in sheets:
            ws = workbook.create_sheet(_safe_sheet_name(sheet.name))
            for column_index, title in enumerate(sheet.columns, start=1):
                ws.cell(row=1, column=column_index, value=title)
            for row_index, row in enumerate(sheet.rows, start=2):
                for column_index, value in enumerate(row, start=1):
                    ws.cell(row=row_index, column=column_index, value=value)
            _autosize_columns(ws)
        normalized_output_root = _safe_output_root(
            output_root
            or _active_output_root(self.ctx)
            or _default_sheet_output_root(self.ctx)
        )
        output_dir = self.ctx.work_dir / normalized_output_root
        output_dir.mkdir(parents=True, exist_ok=True)
        filename = f"{_safe_file_stem(draft.artifact_id or draft.title or 'sheet')}.xlsx"
        path = output_dir / filename
        workbook.save(path)
        return RenderedSheetArtifact(
            path=path,
            relative_path=f"{normalized_output_root}/{filename}",
            metadata={
                "artifact_id": draft.artifact_id,
                "family": str(draft.family),
                "title": draft.title,
                "sheet_count": len(sheets),
            },
        )


def _coerce_sheet(raw: StructuredSheet | dict[str, Any]) -> StructuredSheet:
    if isinstance(raw, StructuredSheet):
        return raw
    return StructuredSheet.model_validate(raw)


def _safe_sheet_name(name: str) -> str:
    safe = re.sub(r"[\[\]:*?/\\]", " ", str(name or "Sheet")).strip() or "Sheet"
    return safe[:31]


def _safe_file_stem(value: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", str(value or "sheet")).strip("-")
    return safe or "sheet"


def _safe_output_root(value: str) -> str:
    normalized = str(value or "").replace("\\", "/").strip().strip("/")
    if normalized.startswith("project/"):
        normalized = normalized.removeprefix("project/").strip("/")
    first_segment = normalized.split("/", 1)[0].strip()
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", first_segment).strip("-")
    return safe or "xlsx-prepared"


def _active_output_root(ctx: "HarnessContext") -> str | None:
    workspace_session = getattr(ctx, "workspace_runtime_session", None)
    if isinstance(workspace_session, dict):
        root = str(workspace_session.get("artifact_work_root") or "").strip()
        if root:
            return root
        agent_cwd = str(workspace_session.get("agent_cwd") or "").strip()
        if agent_cwd.startswith("project/"):
            return agent_cwd.removeprefix("project/").strip("/") or None
    artifact_work_root = str(getattr(ctx, "artifact_work_root", "") or "").strip()
    return artifact_work_root or None


def _default_sheet_output_root(ctx: "HarnessContext") -> str:
    skill_id = str(getattr(ctx, "skill_id", None) or "xlsx").strip()
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", skill_id).strip("-")
    return f"{safe or 'xlsx'}-prepared"


def _autosize_columns(ws) -> None:
    from openpyxl.utils import get_column_letter

    for column_cells in ws.columns:
        max_len = 8
        for cell in column_cells:
            max_len = max(max_len, len(str(cell.value or "")) + 2)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max_len, 48)
